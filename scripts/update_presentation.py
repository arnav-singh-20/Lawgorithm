import pptx
import os

ppt_path = "Lawgorithm_Review2_Presentation.pptx"
prs = pptx.Presentation(ppt_path)

# 1. Slide 3: Abstract (Point 5 / Shape 30)
slide3 = prs.slides[2]
shape30 = slide3.shapes[30]
if shape30.has_text_frame and shape30.text_frame.paragraphs:
    p = shape30.text_frame.paragraphs[0]
    if p.runs:
        p.runs[0].text = "Next: Multi-tool expansion, human verification, formal evaluation, and legal corpus RAG dataset integration."
        for r in p.runs[1:]:
            r.text = ""

# 2. Slide 10: Updated Methodology
slide10 = prs.slides[9]
# Phase 2 header (Shape 12)
p = slide10.shapes[12].text_frame.paragraphs[0]
p.runs[0].text = "Phase 2  —  Tools & Trust Pipeline"
for r in p.runs[1:]:
    r.text = ""

# Phase 2 bullets (Shape 13)
tf13 = slide10.shapes[13].text_frame
tf13.paragraphs[0].runs[0].text = "Multi-tool split: segment, classify risk, translate"
for r in tf13.paragraphs[0].runs[1:]:
    r.text = ""

tf13.paragraphs[1].runs[0].text = "Schema-validated structured JSON output & UI/API"
for r in tf13.paragraphs[1].runs[1:]:
    r.text = ""

tf13.paragraphs[2].runs[0].text = "Human verification dashboard for low-confidence output"
for r in tf13.paragraphs[2].runs[1:]:
    r.text = ""

# Phase 3 header (Shape 16)
p = slide10.shapes[16].text_frame.paragraphs[0]
p.runs[0].text = "Phase 3  —  Legal Corpus RAG & Scale"
for r in p.runs[1:]:
    r.text = ""

# Phase 3 bullets (Shape 17)
tf17 = slide10.shapes[17].text_frame
tf17.paragraphs[0].runs[0].text = "Legal corpus dataset collection & section-level chunking"
for r in tf17.paragraphs[0].runs[1:]:
    r.text = ""

tf17.paragraphs[1].runs[0].text = "Vector retrieval (RAG) integration into ReAct agent loop"
for r in tf17.paragraphs[1].runs[1:]:
    r.text = ""

tf17.paragraphs[2].runs[0].text = "Formal evaluation pipeline (accuracy / grounding / hallucination)"
for r in tf17.paragraphs[2].runs[1:]:
    r.text = ""

# 3. Slide 13: Implementation Progress
slide13 = prs.slides[12]
# Shape 16 (Item 4 text): Multi-tool split
slide13.shapes[16].text_frame.paragraphs[0].runs[0].text = "Multi-tool split: segmentation, risk classifier, translation"
for r in slide13.shapes[16].text_frame.paragraphs[0].runs[1:]:
    r.text = ""

# Shape 19 (Item 5 text): Frontend / API integration & human verification UI
slide13.shapes[19].text_frame.paragraphs[0].runs[0].text = "Frontend & API integration and verification dashboard"
for r in slide13.shapes[19].text_frame.paragraphs[0].runs[1:]:
    r.text = ""

# Shape 22 (Item 6 text): Legal corpus collection
slide13.shapes[22].text_frame.paragraphs[0].runs[0].text = "Legal corpus collection and section-level chunking (Phase 3)"
for r in slide13.shapes[22].text_frame.paragraphs[0].runs[1:]:
    r.text = ""

# Shape 25 (Item 7 text): Vector retrieval RAG
slide13.shapes[25].text_frame.paragraphs[0].runs[0].text = "Vector retrieval (RAG) integration into agent loop (Phase 3)"
for r in slide13.shapes[25].text_frame.paragraphs[0].runs[1:]:
    r.text = ""

# Shape 28 (Item 8 text): Formal eval
slide13.shapes[28].text_frame.paragraphs[0].runs[0].text = "Formal evaluation pipeline & real-user pilot (Phase 3)"
for r in slide13.shapes[28].text_frame.paragraphs[0].runs[1:]:
    r.text = ""

# 4. Slide 15: Preliminary Results (Shape 5)
slide15 = prs.slides[14]
slide15.shapes[5].text_frame.paragraphs[0].runs[0].text = "Evaluation harness is built; first full-corpus run is pending Phase 3 completion (legal RAG dataset integration). Figures below are the defined targets, not final results."
for r in slide15.shapes[5].text_frame.paragraphs[0].runs[1:]:
    r.text = ""

# 5. Slide 17: Remaining Work
slide17 = prs.slides[16]
# Item 1 (Shape 5, 7)
slide17.shapes[5].text_frame.paragraphs[0].runs[0].text = "Multi-tool expansion"
slide17.shapes[5].text_frame.paragraphs[1].runs[0].text = "Segmentation, risk classifier, translation as independent tools"
slide17.shapes[7].text_frame.paragraphs[0].runs[0].text = "Weeks 1–2"

# Item 2 (Shape 8, 10)
slide17.shapes[8].text_frame.paragraphs[0].runs[0].text = "API & frontend integration"
slide17.shapes[8].text_frame.paragraphs[1].runs[0].text = "FastAPI wrapper, UI connection, schema-enforced validation"
slide17.shapes[10].text_frame.paragraphs[0].runs[0].text = "Weeks 2–3"

# Item 3 (Shape 11, 13)
slide17.shapes[11].text_frame.paragraphs[0].runs[0].text = "Human verification dashboard"
slide17.shapes[11].text_frame.paragraphs[1].runs[0].text = "Verifier view: approve / edit / reject, logged corrections"
slide17.shapes[13].text_frame.paragraphs[0].runs[0].text = "Weeks 3–4"

# Item 4 (Shape 14, 16)
slide17.shapes[14].text_frame.paragraphs[0].runs[0].text = "Legal corpus collection & chunking"
slide17.shapes[14].text_frame.paragraphs[1].runs[0].text = "Curating Indian Contract Act, Rent Control & Labour statutes"
slide17.shapes[16].text_frame.paragraphs[0].runs[0].text = "Weeks 4–5"

# Item 5 (Shape 17, 19)
slide17.shapes[17].text_frame.paragraphs[0].runs[0].text = "Vector retrieval (RAG) integration"
slide17.shapes[17].text_frame.paragraphs[1].runs[0].text = "Embeddings, vector store retrieval grounding in agent loop"
slide17.shapes[19].text_frame.paragraphs[0].runs[0].text = "Weeks 5–6"

# Item 6 (Shape 20, 22)
slide17.shapes[20].text_frame.paragraphs[0].runs[0].text = "Formal evaluation & compliance"
slide17.shapes[20].text_frame.paragraphs[1].runs[0].text = "Accuracy, grounding precision, hallucination rate & DPDP"
slide17.shapes[22].text_frame.paragraphs[0].runs[0].text = "Weeks 6–7"

# Item 7 (Shape 23, 25)
slide17.shapes[23].text_frame.paragraphs[0].runs[0].text = "Pilot with real users"
slide17.shapes[23].text_frame.paragraphs[1].runs[0].text = "Law college / NGO / placement cell pilot and feedback"
slide17.shapes[25].text_frame.paragraphs[0].runs[0].text = "Weeks 7–8"

prs.save(ppt_path)
print("Lawgorithm_Review2_Presentation.pptx successfully updated!")
